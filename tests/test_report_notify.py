"""日报增强的测试：平台维度聚合、Excel 汇总 Sheet、Webhook 推送、定时可配置。

为什么要单独一个文件：
`tests/test_report_service.py` 覆盖的是「生成链路 + 幂等 + AI 降级」，
本次新增的三块能力（平台透视、推送、定时配置化）与生成链路是正交的，
分开写便于定位失败。

推送相关用例**全部离线**：把 httpx.AsyncClient 换成假实现，
既验证报文构造，也验证「HTTP 200 但 errcode != 0 必须算失败」这类
最容易骗过监控的假成功。
"""
import json
from io import BytesIO

import pytest
import pytest_asyncio
from openpyxl import load_workbook

from app.config import settings
from app.models.job import Job
from app.services import report_service, webhook_service
from app.utils.job_key import generate_job_key
from app.utils.timeutil import utc_now


# --------------------------------------------------------------------------
# 样本数据
# --------------------------------------------------------------------------


@pytest.fixture
def multi_platform_jobs(db_session):
    """两个平台 × 三个城市的岗位样本（只在架岗位参与统计）。"""
    rows = [
        # (title, city, source_site, status)
        ("Java后端", "北京", "boss", "ACTIVE"),
        ("Python开发", "北京", "boss", "ACTIVE"),
        ("前端工程师", "北京", "51job", "ACTIVE"),
        ("算法工程师", "上海", "boss", "ACTIVE"),
        ("测试工程师", "上海", "51job", "ACTIVE"),
        ("运维工程师", "深圳", "51job", "ACTIVE"),
        ("已下架岗位", "深圳", "boss", "OFFLINE"),   # 不应计入统计
    ]
    for index, (title, city, site, status) in enumerate(rows):
        db_session.add(Job(
            title=title,
            company_name=f"公司{index}",
            city=city,
            source_site=site,
            job_key=generate_job_key(site, title, f"公司{index}", city),
            job_status=status,
            min_salary=15000,
            max_salary=25000,
            experience="1-3年",
            education="本科",
            skills="Python",
            job_desc=title,
            last_seen_at=utc_now(),
        ))
    db_session.commit()


# --------------------------------------------------------------------------
# 平台维度聚合
# --------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_平台分布只统计在架岗位(db_session, multi_platform_jobs):
    stats = await report_service.build_platform_stats(db_session)
    by_name = {item["name"]: item["count"] for item in stats}

    assert by_name == {"51job": 3, "boss": 3}      # OFFLINE 那条不计入
    # 按数量倒序（并列时按名称）
    assert [item["count"] for item in stats] == sorted(
        [item["count"] for item in stats], reverse=True
    )


@pytest.mark.asyncio
async def test_城市平台透视表结构(db_session, multi_platform_jobs):
    pivot = await report_service.build_city_platform_pivot(db_session)

    assert set(pivot["platforms"]) == {"boss", "51job"}
    assert pivot["grand_total"] == 6
    assert pivot["totals"] == {"boss": 3, "51job": 3}

    by_city = {row["city"]: row for row in pivot["rows"]}
    assert by_city["北京"]["counts"] == {"boss": 2, "51job": 1}
    assert by_city["北京"]["total"] == 3
    assert by_city["上海"]["total"] == 2
    assert by_city["深圳"]["total"] == 1

    # 城市按总量倒序
    assert [row["total"] for row in pivot["rows"]] == [3, 2, 1]


@pytest.mark.asyncio
async def test_城市平台透视表带合计行列自洽(db_session, multi_platform_jobs):
    """每行合计必须等于该行各平台之和，总计必须等于各列之和。"""
    pivot = await report_service.build_city_platform_pivot(db_session)

    for row in pivot["rows"]:
        assert row["total"] == sum(row["counts"].values())
    assert pivot["grand_total"] == sum(pivot["totals"].values())
    assert pivot["grand_total"] == sum(row["total"] for row in pivot["rows"])


@pytest.mark.asyncio
async def test_透视表超额城市被截断并如实标注(db_session, multi_platform_jobs):
    pivot = await report_service.build_city_platform_pivot(db_session, top_n=2)

    assert len(pivot["rows"]) == 2
    assert pivot["truncated_cities"] == 1          # 3 个城市只列 2 个
    # 截断的是城市行，总计仍必须是全量，不能只合计被列出的行
    assert pivot["grand_total"] == 6


@pytest.mark.asyncio
async def test_空库时透视表不报错(db_session):
    pivot = await report_service.build_city_platform_pivot(db_session)
    assert pivot["rows"] == []
    assert pivot["grand_total"] == 0


@pytest.mark.asyncio
async def test_日报统计快照包含平台维度(db_session, multi_platform_jobs):
    stats = await report_service.build_report_stats(db_session)
    assert "platform" in stats
    assert "city_platform_pivot" in stats
    # 快照必须可 JSON 序列化（要写进 stats_json 字段）
    json.dumps(stats, ensure_ascii=False)


# --------------------------------------------------------------------------
# Excel 汇总 Sheet
# --------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_excel_包含汇总统计与平台分布_sheet(db_session, multi_platform_jobs):
    stats = await report_service.build_report_stats(db_session)
    workbook = load_workbook(BytesIO(report_service.build_excel(stats)))

    for name in ("总览", "汇总统计", "平台分布", "城市分布", "薪资分布"):
        assert name in workbook.sheetnames, f"缺少 sheet：{name}"


@pytest.mark.asyncio
async def test_excel_透视表写入合计行列(db_session, multi_platform_jobs):
    stats = await report_service.build_report_stats(db_session)
    workbook = load_workbook(BytesIO(report_service.build_excel(stats)))
    ws = workbook["汇总统计"]

    header = [cell.value for cell in ws[1]]
    assert header[0] == "城市"
    assert header[-1] == "合计"
    assert "boss" in header and "51job" in header

    # 最后一行的第一列是合计行标签
    last_row = [cell.value for cell in ws[ws.max_row]]
    assert last_row[0] == "合计"
    assert last_row[-1] == 6


@pytest.mark.asyncio
async def test_excel_平台分布_sheet_有数据行(db_session, multi_platform_jobs):
    stats = await report_service.build_report_stats(db_session)
    workbook = load_workbook(BytesIO(report_service.build_excel(stats)))
    ws = workbook["平台分布"]

    assert [cell.value for cell in ws[1]] == ["来源平台", "岗位数"]
    rows = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(2, ws.max_row + 1)}
    assert rows == {"boss": 3, "51job": 3}


def test_excel_无平台数据时不崩():
    """老快照（本次改动之前生成的）里没有 platform 字段，必须照样能出报表。"""
    stats = {"date": "2026-09-17", "summary": {"total": 1, "active": 1, "avg_salary": 9000}}
    workbook = load_workbook(BytesIO(report_service.build_excel(stats)))
    assert "平台分布" in workbook.sheetnames
    assert "汇总统计" in workbook.sheetnames


# --------------------------------------------------------------------------
# Webhook 推送
# --------------------------------------------------------------------------


class FakeResponse:
    def __init__(self, status_code=200, payload=None, raise_json=False):
        self.status_code = status_code
        self._payload = payload if payload is not None else {"errcode": 0, "errmsg": "ok"}
        self._raise_json = raise_json

    def json(self):
        if self._raise_json:
            raise ValueError("not json")
        return self._payload


@pytest.fixture
def fake_http(monkeypatch):
    """替换 httpx.AsyncClient，记录发出的报文。"""
    calls = []

    def install(status_code=200, payload=None, exc=None, raise_json=False):
        class _Client:
            def __init__(self, *args, **kwargs):
                pass

            async def __aenter__(self):
                return self

            async def __aexit__(self, *exc_info):
                return False

            async def post(self, url, json=None):
                calls.append({"url": url, "json": json})
                if exc is not None:
                    raise exc
                return FakeResponse(status_code, payload, raise_json)

        monkeypatch.setattr(webhook_service.httpx, "AsyncClient", _Client)
        return calls

    return install


@pytest.fixture
def webhook_on(monkeypatch):
    """打开推送并给出一个合法配置。"""
    monkeypatch.setattr(settings, "report_webhook_enabled", True)
    monkeypatch.setattr(settings, "report_webhook_type", "wecom")
    monkeypatch.setattr(settings, "report_webhook_url", "https://example.com/hook?key=abc")


def test_类型归一化支持常见别名():
    assert webhook_service.normalize_type("wecom") == "wecom"
    assert webhook_service.normalize_type("企业微信") == "wecom"
    assert webhook_service.normalize_type("dingtalk") == "dingtalk"
    assert webhook_service.normalize_type("钉钉") == "dingtalk"
    assert webhook_service.normalize_type("generic") == "generic"
    assert webhook_service.normalize_type(None) == "wecom"
    # 拼错的类型不能被静默当成某一类，否则报文会以错误格式发出去
    assert webhook_service.normalize_type("wechat") is None


def test_各类型报文格式():
    wecom = webhook_service.build_payload("wecom", "标题", "正文")
    assert wecom == {"msgtype": "markdown", "markdown": {"content": "正文"}}

    ding = webhook_service.build_payload("dingtalk", "标题", "正文")
    assert ding == {"msgtype": "markdown", "markdown": {"title": "标题", "text": "正文"}}

    generic = webhook_service.build_payload("generic", "标题", "正文")
    assert generic == {"title": "标题", "content": "正文"}


@pytest.mark.asyncio
async def test_未开启推送时不发请求(fake_http, monkeypatch):
    calls = fake_http()
    monkeypatch.setattr(settings, "report_webhook_enabled", False)

    result = await webhook_service.push("标题", "正文")

    assert result.attempted is False
    assert result.success is False
    assert calls == []


@pytest.mark.asyncio
async def test_开启但未配_url_时不发请求(fake_http, monkeypatch):
    calls = fake_http()
    monkeypatch.setattr(settings, "report_webhook_enabled", True)
    monkeypatch.setattr(settings, "report_webhook_url", "   ")

    result = await webhook_service.push("标题", "正文")

    assert result.attempted is False
    assert "report_webhook_url" in result.detail
    assert calls == []


@pytest.mark.asyncio
async def test_类型非法时明确报错而不是发错格式(fake_http, webhook_on, monkeypatch):
    calls = fake_http()
    monkeypatch.setattr(settings, "report_webhook_type", "wechat")

    result = await webhook_service.push("标题", "正文")

    assert result.attempted is False
    assert "wechat" in result.detail
    assert calls == []


@pytest.mark.asyncio
async def test_url_协议非法时拦截(fake_http, webhook_on, monkeypatch):
    calls = fake_http()
    monkeypatch.setattr(settings, "report_webhook_url", "qyapi.weixin.qq.com/hook")

    result = await webhook_service.push("标题", "正文")

    assert result.attempted is False
    assert calls == []


@pytest.mark.asyncio
async def test_推送成功(fake_http, webhook_on):
    calls = fake_http(payload={"errcode": 0, "errmsg": "ok"})

    result = await webhook_service.push("标题", "正文")

    assert result.attempted and result.success
    assert len(calls) == 1
    assert calls[0]["json"]["msgtype"] == "markdown"


@pytest.mark.asyncio
async def test_http200_但_errcode_非零必须算失败(fake_http, webhook_on):
    """企业微信/钉钉用 200 + errcode 表达失败，只看状态码会把失效 token 记成成功。"""
    fake_http(payload={"errcode": 93000, "errmsg": "invalid webhook url"})

    result = await webhook_service.push("标题", "正文")

    assert result.attempted is True
    assert result.success is False
    assert "93000" in result.detail


@pytest.mark.asyncio
async def test_http_非200_算失败(fake_http, webhook_on):
    fake_http(status_code=500, raise_json=True)

    result = await webhook_service.push("标题", "正文")

    assert result.success is False
    assert "500" in result.detail


@pytest.mark.asyncio
async def test_网络异常不得抛出(fake_http, webhook_on):
    """推送失败不能把日报生成拖崩 —— 必须返回失败结果而不是抛异常。"""
    fake_http(exc=RuntimeError("connection reset"))

    result = await webhook_service.push("标题", "正文")

    assert result.attempted is True
    assert result.success is False
    assert "connection reset" in result.detail


@pytest.mark.asyncio
async def test_is_configured_口径(monkeypatch):
    monkeypatch.setattr(settings, "report_webhook_enabled", False)
    monkeypatch.setattr(settings, "report_webhook_url", "https://example.com/hook")
    monkeypatch.setattr(settings, "report_webhook_type", "wecom")
    assert webhook_service.is_configured() is False

    monkeypatch.setattr(settings, "report_webhook_enabled", True)
    assert webhook_service.is_configured() is True

    monkeypatch.setattr(settings, "report_webhook_type", "wechat")
    assert webhook_service.is_configured() is False


def test_推送内容包含关键指标并截断长摘要():
    stats = {
        "date": "2026-09-17",
        "new_today": 5,
        "summary": {"total": 100, "active": 90, "avg_salary": 16000},
        "city": [{"name": "北京", "count": 30}, {"name": "上海", "count": 20},
                 {"name": "深圳", "count": 10}],
        "platform": [{"name": "boss", "count": 60}, {"name": "51job", "count": 40}],
        "skill": [{"name": "Python", "count": 50}],
    }

    class _Report:
        report_date = "2026-09-17"
        ai_summary = "很长的摘要" * 300
        generated_by = "rule"

    title, content = report_service.build_webhook_content(_Report(), stats, top_n=2)

    assert title == "招聘市场日报 2026-09-17"
    assert "今日新增：5" in content
    assert "北京(30)" in content
    assert "深圳(10)" not in content          # top_n=2 应截掉
    assert "boss(60)" in content
    # 超长摘要被截断，避免机器人拒收整条消息
    assert len(content) < 1500


@pytest.mark.asyncio
async def test_push_report_未配置时安静跳过(db_session, monkeypatch):
    monkeypatch.setattr(settings, "report_webhook_enabled", False)

    class _Report:
        report_date = "2026-09-17"
        ai_summary = "摘要"
        generated_by = "rule"

    result = await report_service.push_report(_Report(), {})
    assert result.attempted is False


@pytest.mark.asyncio
async def test_push_report_内部异常被吞掉(monkeypatch):
    """即使构造内容这一步就炸了，也不能把异常抛给调用方。"""
    class _Boom:
        report_date = "2026-09-17"

        @property
        def ai_summary(self):
            raise RuntimeError("boom")

    result = await report_service.push_report(_Boom(), {"summary": {}})
    assert result.attempted is False
    assert "boom" in result.detail


# --------------------------------------------------------------------------
# 定时任务配置化
# --------------------------------------------------------------------------


def test_时间参数越界或非法时回退默认值():
    from app.scheduler import _clamp

    assert _clamp(3, 0, 23, 2) == 3
    assert _clamp("7", 0, 23, 2) == 7
    assert _clamp(99, 0, 23, 2) == 2
    assert _clamp(-1, 0, 23, 2) == 2
    assert _clamp("abc", 0, 23, 2) == 2
    assert _clamp(None, 0, 23, 2) == 2


@pytest.fixture
def clean_scheduler_job():
    """确保测试用的任务在结束后从全局调度器里摘掉。"""
    from app.scheduler import scheduler

    def cleanup(job_id):
        from app.scheduler import scheduler as sched
        if sched.get_job(job_id):
            sched.remove_job(job_id)

    yield cleanup

    for job_id in ("__test_job__", "__test_disabled__"):
        if scheduler.get_job(job_id):
            scheduler.remove_job(job_id)


def test_开关关闭时不注册任务(clean_scheduler_job):
    from app.scheduler import _sync_job

    async def _noop():
        pass

    _sync_job("__test_disabled__", _noop, enabled=False, hour=3, minute=0, label="测试")
    clean_scheduler_job("__test_disabled__")

    from app.scheduler import scheduler
    assert scheduler.get_job("__test_disabled__") is None


def test_触发时间来自配置而不是写死(clean_scheduler_job):
    from app.scheduler import _sync_job, scheduler

    async def _noop():
        pass

    _sync_job("__test_job__", _noop, enabled=True, hour=3, minute=17, label="测试")

    job = scheduler.get_job("__test_job__")
    assert job is not None
    fields = {field.name: str(field) for field in job.trigger.fields}
    assert fields["hour"] == "3"
    assert fields["minute"] == "17"

    clean_scheduler_job("__test_job__")


def test_重复注册不会产生重复任务(clean_scheduler_job):
    from app.scheduler import _sync_job, scheduler

    async def _noop():
        pass

    _sync_job("__test_job__", _noop, enabled=True, hour=1, minute=0, label="测试")
    _sync_job("__test_job__", _noop, enabled=True, hour=5, minute=30, label="测试")

    same_id = [job for job in scheduler.get_jobs() if job.id == "__test_job__"]
    assert len(same_id) == 1, "重复注册必须覆盖而不是新增"

    fields = {field.name: str(field) for field in same_id[0].trigger.fields}
    assert fields["hour"] == "5"    # 以最后一次配置为准
    assert fields["minute"] == "30"

    clean_scheduler_job("__test_job__")


# --------------------------------------------------------------------------
# 推送接口
# --------------------------------------------------------------------------


@pytest_asyncio.fixture
async def admin_client(client, db_session):
    from tests.helpers import admin_token, auth_headers

    token = await admin_token(client, db_session, "admin_reportpush")
    client.headers.update(auth_headers(token))
    return client


async def _seed(db_session):
    db_session.add(Job(
        title="Java开发", company_name="某公司", city="长沙", source_site="boss",
        job_key=generate_job_key("boss", "Java开发", "某公司", "长沙"),
        job_status="ACTIVE", min_salary=15000, max_salary=25000,
        experience="1-3年", education="本科", skills="Java",
        job_desc="Java", last_seen_at=utc_now(),
    ))
    await db_session.commit()


@pytest.mark.asyncio
async def test_手动推送未配置时返回明确原因(admin_client, db_session, monkeypatch):
    monkeypatch.setattr(settings, "report_webhook_enabled", False)
    await _seed(db_session)

    generated = await admin_client.post("/api/reports/generate")
    report_id = generated.json()["data"]["id"]

    resp = await admin_client.post(f"/api/reports/{report_id}/push")

    assert resp.status_code == 400
    assert "未执行推送" in resp.json()["message"]


@pytest.mark.asyncio
async def test_手动推送成功(admin_client, db_session, monkeypatch, fake_http):
    fake_http(payload={"errcode": 0, "errmsg": "ok"})
    monkeypatch.setattr(settings, "report_webhook_enabled", True)
    monkeypatch.setattr(settings, "report_webhook_type", "wecom")
    monkeypatch.setattr(settings, "report_webhook_url", "https://example.com/hook")
    await _seed(db_session)

    generated = await admin_client.post("/api/reports/generate")
    report_id = generated.json()["data"]["id"]

    resp = await admin_client.post(f"/api/reports/{report_id}/push")

    assert resp.status_code == 200, resp.text
    data = resp.json()["data"]
    assert data["success"] is True
    assert data["target"] == "企业微信机器人"


@pytest.mark.asyncio
async def test_推送不存在的日报返回404(admin_client):
    resp = await admin_client.post("/api/reports/99999/push")
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_详情接口透出推送配置状态(admin_client, db_session, monkeypatch):
    monkeypatch.setattr(settings, "report_webhook_enabled", False)
    await _seed(db_session)
    generated = await admin_client.post("/api/reports/generate")
    report_id = generated.json()["data"]["id"]

    detail = await admin_client.get(f"/api/reports/{report_id}")

    push = detail.json()["data"]["push"]
    assert push["enabled"] is False
    assert push["configured"] is False
