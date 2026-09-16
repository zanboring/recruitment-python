"""自动化日报模块测试：幂等、统计、Excel、AI 摘要降级、API 链路。"""
import json
import pytest

from app.models.job import Job
from app.services import report_service
from app.utils.timeutil import utc_now


@pytest.fixture
def sample_jobs(db_session):
    """插入少量岗位样本，供日报统计使用。"""
    from app.utils.job_key import generate_job_key

    rows = [
        ("Java后端工程师", "字节跳动", "北京", 20000, 35000, "3-5年", "本科", "Java,Spring"),
        ("Python开发", "腾讯", "深圳", 18000, 30000, "1-3年", "本科", "Python,FastAPI"),
        ("前端工程师", "美团", "北京", 15000, 25000, "1-3年", "本科", "Vue,TypeScript"),
        ("算法工程师", "字节跳动", "上海", 30000, 50000, "3-5年", "硕士", "Python,机器学习"),
    ]
    jobs = []
    for title, company, city, lo, hi, exp, edu, skills in rows:
        job = Job(
            title=title,
            company_name=company,
            city=city,
            source_site="boss",
            job_key=generate_job_key("boss", title, company, city),
            job_status="ACTIVE",
            min_salary=lo,
            max_salary=hi,
            experience=exp,
            education=edu,
            skills=skills,
            job_desc=title,
            last_seen_at=utc_now(),
        )
        db_session.add(job)
        jobs.append(job)
    db_session.commit()
    return jobs


@pytest.mark.asyncio
async def test_build_report_stats_dimensions(db_session, sample_jobs):
    """统计快照应覆盖日期/新增/汇总/城市/技能等关键维度。"""
    stats = await report_service.build_report_stats(db_session)
    assert stats["date"]
    assert stats["new_today"] == len(sample_jobs)
    assert stats["summary"]["active"] == len(sample_jobs)
    assert stats["summary"]["total"] == len(sample_jobs)
    assert len(stats["city"]) >= 2
    city_names = {c["name"] for c in stats["city"]}
    assert "北京" in city_names and "上海" in city_names


@pytest.mark.asyncio
async def test_build_excel_contains_sheets(db_session, sample_jobs):
    """Excel 应包含总览 + 多个维度 sheet。"""
    stats = await report_service.build_report_stats(db_session)
    data = report_service.build_excel(stats)
    assert data[:2] == b"PK"  # xlsx 是 zip 包
    from openpyxl import load_workbook
    from io import BytesIO

    wb = load_workbook(filename=BytesIO(data))
    names = wb.sheetnames
    for expect in ["总览", "城市分布", "热门技能", "薪资分布"]:
        assert expect in names


@pytest.mark.asyncio
async def test_create_report_idempotent(db_session, sample_jobs, monkeypatch):
    """同一天重复生成应返回同一条记录（幂等），不产生重复日报。"""
    async def fake_summarize(stats, db):
        return "规则摘要", "rule"

    monkeypatch.setattr(report_service, "summarize_with_ai", fake_summarize)

    first = await report_service.create_daily_report(db_session, None)
    second = await report_service.create_daily_report(db_session, None)
    assert first.id == second.id
    assert first.status == "GENERATED"
    assert first.report_date == second.report_date


@pytest.mark.asyncio
async def test_create_report_ai_summary_fallback(db_session, sample_jobs, monkeypatch):
    """AI 摘要失败时应回退到规则文本，且生成层标记为 rule。"""
    async def broken_generate(stats, db):
        raise RuntimeError("LLM 全挂")

    import app.services.ai_service as ai_service

    monkeypatch.setattr(ai_service, "generate_analysis_report", broken_generate)
    report = await report_service.create_daily_report(db_session, None)
    assert report.status == "GENERATED"
    assert report.generated_by == "rule"
    assert "新增" in report.ai_summary


@pytest.mark.asyncio
async def test_report_dict_roundtrip(db_session, sample_jobs, monkeypatch):
    """report_to_dict / detail 应还原统计快照。"""
    async def fake_summarize(stats, db):
        return "AI摘要", "primary"

    monkeypatch.setattr(report_service, "summarize_with_ai", fake_summarize)
    report = await report_service.create_daily_report(db_session, None)
    detail = report_service.report_detail_to_dict(report)
    assert detail["generated_by"] == "primary"
    assert detail["new_jobs"] == len(sample_jobs)
    assert detail["stats"]["summary"]["active"] == len(sample_jobs)