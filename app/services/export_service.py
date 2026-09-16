"""岗位数据 Excel 导入导出。

设计要点：
- **导入去重口径与爬虫、管理端新增完全一致**（app.utils.job_key）。原先导入用的
  job_key 含行号与整行内容哈希，同一份文件二次导入会生成不同的键，结果要么
  重复入库、要么撞唯一约束导致整个事务回滚、一条都进不去。
- 已存在的记录按「跳过」处理并计入 skip，而不是报错回滚。
- 导出用 model_copy 复制 DTO：直接改传入对象会污染调用方状态（原实现有此副作用）。
- 导出与导入都设了体量上限，避免一次操作把内存打满。
"""
import logging
from io import BytesIO
from datetime import datetime, timezone

from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.job import Job
from app.services.job_service import JobService
from app.schemas.job import JobQueryDTO
from app.utils.job_key import generate_job_key

logger = logging.getLogger(__name__)

# 单次导出行数上限：避免误点导出把几十万行一次性读进内存
EXPORT_MAX_ROWS = 50000

# 单次导入文件大小上限（字节）。xlsx 是压缩包，10MB 已可容纳数十万行
IMPORT_MAX_BYTES = 10 * 1024 * 1024

EXPORT_HEADERS = [
    "标题", "公司", "城市", "薪资(min)", "薪资(max)", "经验",
    "学历", "技能", "来源平台", "发布时间", "描述",
]


def _to_float(value) -> float:
    """把单元格值安全转成 float；空值或脏数据返回 0.0。

    不用 float(...) 裸转：单元格为空时 float(None) 会抛 TypeError，
    整行会被判为失败行。
    """
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


async def export_jobs_to_excel(db: AsyncSession, query_dto: JobQueryDTO):
    """按查询条件导出岗位为 xlsx 文件流。"""
    # 复制一份再改分页：直接改传入的 DTO 会污染调用方对象状态
    export_dto = query_dto.model_copy(update={"page_num": 1, "page_size": EXPORT_MAX_ROWS})
    jobs = await JobService.query_jobs(db, export_dto)

    wb = Workbook()
    ws = wb.active
    ws.title = "岗位数据"
    ws.append(EXPORT_HEADERS)

    for job in jobs:
        publish_time = job.publish_time.strftime("%Y-%m-%d") if job.publish_time else ""
        ws.append([
            job.title,
            job.company_name,
            job.city,
            job.min_salary,
            job.max_salary,
            job.experience,
            job.education,
            job.skills,
            job.source_site,
            publish_time,
            job.job_desc or "",
        ])

    for col in ws.columns:
        max_length = 0
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except (TypeError, AttributeError):
                pass
        adjusted_width = (max_length + 2) * 1.2
        ws.column_dimensions[col[0].column_letter].width = min(adjusted_width, 50)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"jobs_export_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )



async def import_jobs_from_csv(db: AsyncSession, file_content: bytes) -> dict:
    """从 CSV 导入岗位，返回 {"success", "skip", "fail"} 统计（幂等）。

    表头字段与 xlsx 导出一致（中文表头）：
    标题,公司,城市,薪资(min),薪资(max),经验,学历,技能,来源平台,发布时间,描述

    复用 import_jobs_from_excel 的逐行容错与 job_key 去重口径 ——
    同一岗位从 CSV / 爬虫 / 管理端三个入口进入都不会重复入库。
    """
    import csv
    import io

    if not file_content:
        raise ValueError("文件内容为空")
    if len(file_content) > IMPORT_MAX_BYTES:
        raise ValueError(f"文件过大，上限 {IMPORT_MAX_BYTES // 1024 // 1024}MB")

    success_count = 0
    skip_count = 0
    fail_count = 0

    try:
        text = file_content.decode("utf-8-sig")  # 兼容 UTF-8 BOM（Excel 导出常见）
        reader = csv.DictReader(io.StringIO(text))

        for raw in reader:
            try:
                title = str(raw.get("标题") or "").strip()
                if not title:
                    continue

                company_name = str(raw.get("公司") or "").strip()
                city = str(raw.get("城市") or "").strip()
                source_site = str(raw.get("来源平台") or "").strip() or "import"
                job_key = generate_job_key(source_site, title, company_name, city)

                existing = await db.execute(
                    select(Job.id).where(Job.job_key == job_key)
                )
                if existing.scalar_one_or_none() is not None:
                    skip_count += 1
                    continue

                db.add(Job(
                    title=title,
                    company_name=company_name,
                    city=city,
                    min_salary=_to_float(raw.get("薪资(min)")),
                    max_salary=_to_float(raw.get("薪资(max)")),
                    experience=str(raw.get("经验") or ""),
                    education=str(raw.get("学历") or ""),
                    skills=str(raw.get("技能") or ""),
                    source_site=source_site,
                    job_key=job_key,
                    job_status="ACTIVE",
                    job_desc=str(raw.get("描述") or ""),
                ))
                success_count += 1
            except Exception as e:  # noqa: BLE001 - 单行脏数据不影响其余行
                fail_count += 1
                logger.warning("CSV 第 %s 行导入失败：%s", raw, e)

        await db.commit()
    except Exception as e:
        await db.rollback()
        raise ValueError(f"CSV 导入失败: {e}")

    return {"success": success_count, "skip": skip_count, "fail": fail_count}


async def import_jobs_from_excel(db: AsyncSession, file_content: bytes) -> dict:
    """从 xlsx 导入岗位，返回 {"success", "skip", "fail"} 统计。

    - job_key 相同的记录计入 skip（幂等：同一份文件重复导入不会产生重复数据，
      也不会因唯一约束冲突而整体回滚）；
    - 逐行独立 try/except，单行脏数据不影响其余行；
    - 空行（Excel 常见的尾部空行）直接忽略，不计入失败。
    """
    if not file_content:
        raise ValueError("文件内容为空")
    if len(file_content) > IMPORT_MAX_BYTES:
        raise ValueError(f"文件过大，上限 {IMPORT_MAX_BYTES // 1024 // 1024}MB")

    success_count = 0
    skip_count = 0
    fail_count = 0

    try:
        wb = load_workbook(filename=BytesIO(file_content), read_only=True)
        ws = wb.active

        rows = ws.iter_rows(values_only=True)
        raw_headers = next(rows, None)
        if raw_headers is None:
            raise ValueError("文件没有表头")
        headers = [str(h) if h is not None else "" for h in raw_headers]

        for values in rows:
            try:
                raw = dict(zip(headers, values))

                title = str(raw.get("标题") or "").strip()
                if not title:
                    continue  # 空行不计入失败

                company_name = str(raw.get("公司") or "").strip()
                city = str(raw.get("城市") or "").strip()
                source_site = str(raw.get("来源平台") or "").strip() or "import"

                # 与爬虫 / 管理端新增共用同一套指纹规则
                job_key = generate_job_key(source_site, title, company_name, city)

                existing = await db.execute(select(Job.id).where(Job.job_key == job_key))
                if existing.scalar_one_or_none() is not None:
                    skip_count += 1
                    continue

                db.add(Job(
                    title=title,
                    company_name=company_name,
                    city=city,
                    min_salary=_to_float(raw.get("薪资(min)")),
                    max_salary=_to_float(raw.get("薪资(max)")),
                    experience=str(raw.get("经验") or ""),
                    education=str(raw.get("学历") or ""),
                    skills=str(raw.get("技能") or ""),
                    source_site=source_site,
                    job_key=job_key,
                    job_status="ACTIVE",
                    job_desc=str(raw.get("描述") or ""),
                ))
                success_count += 1
            except Exception as e:  # noqa: BLE001 - 单行脏数据不影响其余行
                fail_count += 1
                logger.warning("导入第 %s 行失败：%s", values, e)

        await db.commit()
    except Exception as e:
        await db.rollback()
        raise ValueError(f"导入失败: {e}")

    return {"success": success_count, "skip": skip_count, "fail": fail_count}
